import json
import uuid
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from openpyxl import Workbook

from ras_engine import (CATEGORIES, LINKED_FIELDS, BiofilterParams, biofilter_table, link_to_simulation,
                        max_retention_min, simulate, summary, weekly)
from ras_engine.equipment import SIZED_CATEGORIES, normalize_steps
from ras_engine.process import NODE_TYPES

from .forms import (BiofilterForm, FeedForm, ParamsForm, ScenarioForm, StandardForm, same_value,
                    scenario_base_values, stage_field_specs)
from .services import plan_result, prewarm, run_production as _run_production  # _run_production: uncached, for tests
from .parse import _finite, _float_or_none, _int_or_none, _num
from .costing import (_capex_data, _diagram_for, _diagram_lines, _equipment_units, _has_diagram,
                      _line, _resolve, _with_basis)
from .models import HORIZON_UNITS, MAX_HORIZON_DAYS, EquipmentStandard, Feed, ProcessTemplate, Scenario

CHART_KEYS = ["weight_g", "biomass_t", "feed_kg", "volume_m3", "density",
              "q_o2_m3h", "q_tan_m3h", "q_co2_m3h", "q_tss_m3h", "o2_kgd", "q_no3_m3d"]

EXPORT_COLUMNS = [
    ("day", "Day"), ("temperature", "Temperature °C"), ("length_cm", "Length cm"),
    ("weight_g", "Weight g"), ("count", "Fish"), ("biomass_kg", "Biomass kg"),
    ("density", "Density kg/m³"), ("volume_m3", "Required volume m³"), ("feed_name", "Feed"),
    ("protein", "Protein"), ("fcr", "FCR"), ("feed_kg", "Daily feed kg"),
    ("feed_rate", "Feeding rate"), ("feed_cum_t", "Cumulative feed t"), ("sgr", "SGR"),
    ("o2_kgd", "O₂ kg/day"), ("o2_lpm_actual", "O₂ Nl/min (0 °C)"), ("co2_kgd", "CO₂ kg/day"),
    ("ozone_dose", "Ozone dose g/kg feed"), ("ozone_kgd", "Ozone kg/day"),
    ("ozone_lpm", "Ozone Nl/min (0 °C, pure O₃)"), ("o2_sat_in", "O₂ saturation inlet"),
    ("o2_sat_out", "O₂ saturation outlet"), ("q_o2_m3h", "Q O₂ m³/h"),
    ("q_tan_m3h", "Q TAN m³/h"), ("q_co2_m3h", "Q CO₂ m³/h"), ("q_tss_m3h", "Q TSS m³/h"),
    ("q_design_m3h", "Design flow m³/h"), ("q_no3_m3d", "NO₃ water exchange m³/day"),
    ("water_per_feed", "Water L/kg feed"), ("backflush60_m3d", "Backflush 60 µm m³/day"),
    ("backflush30_m3d", "Backflush 30 µm m³/day"),
]


def _current(request, pk=None):
    if pk:
        scenario = get_object_or_404(Scenario, pk=pk)
    else:
        scenario = Scenario.objects.first()
        if scenario is None:
            scenario = Scenario.objects.create(name="Trout (Excel defaults)")
    from .schema import CURRENT, upgrade  # lazy import, like _media_basis' (schema stays off module import)
    if scenario.schema_version != CURRENT:
        import logging
        for note in upgrade(scenario):  # stamps the row; blobs are never touched (DG-1)
            logging.getLogger("planner").warning("scenario %s (%s): %s", scenario.pk, scenario.name, note)
    return scenario


def _run(scenario):
    feeds = [f.to_engine() for f in Feed.objects.all()] if scenario.use_feed_catalog else None
    return simulate(scenario.engine_params(), feeds=feeds,
                    temperature_steps=scenario.temperature_steps)


def _linked_biofilter(scenario, rows):
    bp, info = link_to_simulation(scenario.engine_biofilter(), rows)
    return bp, info, biofilter_table(bp, scenario.engine_params().design)


def _unique_notes(rows):
    """The rows' advisory safety notes (min. residence, TAN knee), deduplicated —
    the TAN-knee note is identical in every row of the table."""
    seen, out = set(), []
    for r in rows:
        for n in r["notes"]:
            if n not in seen:
                seen.add(n)
                out.append(n)
    return out


def home(request):
    """Start of the workflow: step 01 of the latest scenario (made if there is none yet)."""
    return redirect("brief", pk=_current(request).pk)


def dashboard(request, pk=None):
    scenario = _current(request, pk)
    try:
        rows = _run(scenario)
    except ValueError as exc:  # e.g. the feed table does not cover the weight range
        messages.error(request, str(exc))
        return render(request, "planner/dashboard.html",
                      {"scenario": scenario, "scenarios": Scenario.objects.all(), "error": str(exc)})
    s = summary(rows)
    bio_bp, bio_info, bio_rows = _linked_biofilter(scenario, rows)
    chart = {k: [_finite(r[k]) for r in rows] for k in CHART_KEYS}
    chart["day"] = [r["day"] for r in rows]
    return render(request, "planner/dashboard.html", {
        "scenario": scenario, "scenarios": Scenario.objects.all(), "summary": s,
        "weeks": weekly(rows), "chart": chart, "feeds_used": [
            (name, kg / 1000) for name, kg in s.get("feed_by_type", {}).items()],
        "catalog_empty": not Feed.objects.exists(),
        "bio_bp": bio_bp, "bio_info": bio_info, "bio_rows": bio_rows,
        "bio_notes": _unique_notes(bio_rows),
        "bio_negative": not all(r["feasible"] for r in bio_rows),
        "bio_design_retention": ((bio_info or {}).get("max_retention_min") or 0.0) * bio_bp.max_removal_pct / 100,
    })


def biofilter_design(request, pk):
    """The 'design weight' box on the dashboard."""
    scenario = get_object_or_404(Scenario, pk=pk)
    if request.method == "POST":
        try:
            weight = max(float(request.POST.get("design_weight_g", "0").replace(",", ".")), 0.0)
        except ValueError:
            messages.error(request, "Enter the design weight as a number, e.g. 5.")
        else:
            scenario.biofilter = {**scenario.biofilter, "design_weight_g": weight}
            # biofilter-constants-only save: the plan cache key and ordering intentionally unaffected —
            # the plan forces design_weight_g to 0, and the dashboard's linked biofilter is recomputed
            # from the scenario on every request
            scenario.save(update_fields=["biofilter"])
    return redirect(reverse("dashboard", args=[pk]) + "#biofilter")


def scenario_edit(request, pk=None):
    scenario = get_object_or_404(Scenario, pk=pk) if pk else Scenario(name="New scenario")
    if request.method == "POST":
        meta = ScenarioForm(request.POST, instance=scenario)
        params = ParamsForm(request.POST, values=scenario.params)
        if meta.is_valid() and params.is_valid():
            obj_is_new = scenario.pk is None
            obj = meta.save(commit=False)
            obj.params = params.cleaned_data
            obj.save()
            prewarm(obj)
            messages.success(request, "Scenario saved.")
            # a new project starts at step 01
            return redirect("brief" if obj_is_new else "scenario_edit", pk=obj.pk)
    else:
        meta = ScenarioForm(instance=scenario)
        params = ParamsForm(values=scenario.params)
    g = scenario.engine_params().growth   # server-side: the warning follows what is actually saved
    return render(request, "planner/scenario_form.html",
                  {"scenario": scenario, "scenarios": Scenario.objects.all(),
                   "meta": meta, "params": params,
                   "tgc_warn": g.use_tgc and g.tgc > 2.0})


def scenario_copy(request, pk):
    src = get_object_or_404(Scenario, pk=pk)
    if request.method == "POST":
        src.pk = None
        src.name = f"{src.name} (copy)"
        src.save()
        messages.success(request, "Scenario duplicated. Rename it and adjust the constants.")
        return redirect("scenario_edit", pk=src.pk)
    return redirect("dashboard", pk=pk)


def _snapshot_dir() -> Path:
    """Where delete backups go (tests point this at a temporary directory)."""
    return Path(settings.BASE_DIR) / "backups"


def _backup_json(name: str, payload: dict) -> Path:
    """A full JSON copy written just before a delete, so nothing is lost unrecoverably."""
    directory = _snapshot_dir()
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{name}_deleted_{datetime.now():%Y%m%d-%H%M%S}.json"
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return path


def _remove_process_snapshots(pk):
    """The cached process-drawing PNGs of a scenario go with it."""
    from . import views_docs  # lazy: views_docs imports this module
    for old in views_docs.SNAPSHOT_DIR.glob(f"{pk}-*.png"):
        old.unlink(missing_ok=True)


def scenario_delete(request, pk):
    scenario = get_object_or_404(Scenario, pk=pk)
    if request.method == "POST":
        _backup_json(f"scenario_{scenario.pk}", {
            "name": scenario.name, "schema_version": scenario.schema_version,
            "temperature_steps": scenario.temperature_steps, "use_feed_catalog": scenario.use_feed_catalog,
            "params": scenario.params, "biofilter": scenario.biofilter, "production": scenario.production,
            "equipment": scenario.equipment, "capex": scenario.capex, "process": scenario.process,
            "opex": scenario.opex, "brief": scenario.brief})
        _remove_process_snapshots(pk)
        scenario.delete()
        messages.success(request, "Scenario deleted.")
        if not Scenario.objects.exists():
            return redirect("scenario_new")  # '/' would silently make a new scenario instead
    return redirect("dashboard_home")


def biofilter(request, pk):
    scenario = get_object_or_404(Scenario, pk=pk)
    try:
        sim = _run(scenario)
    except ValueError as exc:  # e.g. the feed table does not cover the weight range
        messages.error(request, str(exc))
        return render(request, "planner/biofilter.html",
                      {"scenario": scenario, "scenarios": Scenario.objects.all(), "error": str(exc)})
    peak_feed = summary(sim).get("peak_feed_kg", 0)
    bp, info, rows = _linked_biofilter(scenario, sim)
    linked = {k: round(getattr(bp, k), 4) for k in LINKED_FIELDS} if info else None
    if request.method == "POST":
        form = BiofilterForm(request.POST, values=scenario.biofilter, linked=linked)
        if form.is_valid():
            data = form.cleaned_data
            if linked:  # fields fed by the simulation: keep the previous manual values
                for k in LINKED_FIELDS:
                    data[k] = scenario.biofilter.get(k, getattr(BiofilterParams(), k))
            if "use_peak" in request.POST:
                data["design_weight_g"] = 0.0
                data["daily_feed_kg"] = round(
                    peak_feed / max(data["n_ras_units"], 1) * data["feed_safety_factor"], 2)
            scenario.biofilter = data
            scenario.save()
            messages.success(request, "Biofilter recalculated.")
            return redirect("biofilter", pk=pk)
    else:
        form = BiofilterForm(values=scenario.biofilter, linked=linked)
    return render(request, "planner/biofilter.html", {
        "scenario": scenario, "scenarios": Scenario.objects.all(), "form": form, "rows": rows,
        "peak_feed": peak_feed, "bp": bp, "info": info, "max_retention": max_retention_min(bp),
        "notes": _unique_notes(rows),
        # what the plan actually designs against (production.py sizes t_bio the same way)
        "design_retention": max_retention_min(bp) * bp.max_removal_pct / 100,
        "negative": not all(r["feasible"] for r in rows),
        "chart": {"t": [r["retention_min"] for r in rows],
                  "salt": [_finite(r["volume_salt_m3"], 3) for r in rows],
                  "fresh": [_finite(r["volume_fresh_m3"], 3) for r in rows]},
    })


def _parse_overrides(post, row: int, base: dict, errors: list) -> dict:
    """Collect the values in a RAS row's constants panel that differ from the scenario."""
    out = {}
    for key, label, _unit, _group, default in stage_field_specs():
        values = post.getlist("ov_" + key)
        if row >= len(values) or values[row] == "":
            continue
        raw = values[row]
        try:
            if isinstance(default, bool):
                value = raw == "True"
            elif isinstance(default, str):
                value = raw
            else:
                value = type(default)(float(raw.replace(",", ".")))
        except ValueError:
            errors.append(f"RAS unit {row + 1} · {label}: enter a number.")
            continue
        if not same_value(value, base.get(key)):
            out[key] = value
    return out


def _parse_stocking(post) -> tuple[dict | None, list[str]]:
    """First stocking date, interval and horizon (step 02)."""
    errors = []
    try:
        start = date.fromisoformat(post.get("start_date", ""))
    except ValueError:
        errors.append("Enter a first stocking date.")
        start = None
    unit = post.get("horizon_unit", "days")
    try:
        interval = int(post.get("interval_days", 0))
        if "horizon_value" in post:
            value = float(post.get("horizon_value", "0").replace(",", "."))
            horizon = round(value * HORIZON_UNITS[unit])
        else:  # plain day count
            value, unit = int(post.get("horizon_days", 0)), "days"
            horizon = value
        if interval < 1 or not 30 <= horizon <= MAX_HORIZON_DAYS:
            raise ValueError
    except (ValueError, KeyError):
        errors.append("Stocking interval must be at least 1 day, and the planning horizon "
                      "between 30 days and 3 years (e.g. 26 weeks, 18 months or 3 years).")
        interval = horizon = value = None
    if errors:
        return None, errors
    return {"start_date": start.isoformat(), "interval_days": interval, "horizon_days": horizon,
            "horizon_value": value if value != int(value) else int(value), "horizon_unit": unit}, []


def _parse_stages(post, base: dict) -> tuple[list | None, list[str]]:
    """RAS units with their ids, exit weights, tanks, retention and constant overrides (step 03)."""
    errors = []
    stages = []
    names = post.getlist("stage_name")
    ids = post.getlist("stage_id") + [""] * len(names)
    rows = zip(names, post.getlist("stage_max"), post.getlist("stage_tanks"), post.getlist("stage_ret"), ids)
    seen = set()
    for i, (name, mx, tanks, ret, uid) in enumerate(rows):
        if not name.strip() and not mx.strip():
            continue
        if not uid or uid in seen:  # a new unit (or a copied row) gets its own id
            uid = uuid.uuid4().hex[:8]
        seen.add(uid)
        try:
            weight, n_tanks, retention = float(mx.replace(",", ".")), int(tanks), float(ret.replace(",", "."))
        except ValueError:
            errors.append(f"RAS unit '{name}': exit weight, tanks and retention time must be numbers.")
            continue
        if n_tanks < 1:
            errors.append(f"RAS unit '{name}': the number of tanks must be at least 1.")
        if retention <= 0:
            errors.append(f"RAS unit '{name}': the retention time must be above zero minutes.")
        stages.append({"id": uid, "name": name.strip() or f"RAS-{len(stages) + 1}",
                       "max_weight_g": weight, "n_tanks": n_tanks, "retention_min": retention,
                       "overrides": _parse_overrides(post, i, base, errors)})
    if not stages:
        errors.append("Add at least one RAS unit.")
    weights = [s["max_weight_g"] for s in stages]
    if any(b <= a for a, b in zip(weights, weights[1:])) or (weights and weights[0] <= 0):
        errors.append("Exit weights must be above zero and increase from one RAS unit to the next.")
    return (None if errors else stages), errors


def _parse_plan(post, base: dict) -> tuple[dict | None, list[str]]:
    stocking, errors = _parse_stocking(post)
    stages, stage_errors = _parse_stages(post, base)
    errors += stage_errors
    if errors:
        return None, errors
    return {**stocking, "stages": stages}, []


def _stage_panels(settings, base):
    """Constants panel per RAS unit: grouped fields, their value and whether they override."""
    specs = stage_field_specs()
    panels = []
    for st in settings["stages"]:
        ov = st.get("overrides") or {}
        groups = {}
        for key, label, unit, group, default in specs:
            value = ov.get(key, base.get(key, default))
            groups.setdefault(group, []).append({
                # full-precision text, so an unchanged value is never saved as an override
                "key": key, "label": label, "unit": unit, "value": value, "text": str(value),
                "is_bool": isinstance(default, bool), "is_text": isinstance(default, str),
                "changed": key in ov, "base": base.get(key, default)})
        panels.append({"stage": st, "groups": list(groups.items()), "n_changed": len(ov)})
    return panels


def production(request, pk):
    scenario = get_object_or_404(Scenario, pk=pk)
    settings = scenario.plan_settings()
    base = scenario_base_values(scenario)
    if request.method == "POST":
        if "stage_name" in request.POST:
            data, errors = _parse_plan(request.POST, base)
        else:
            data, errors = _parse_stocking(request.POST)
            if data:
                data["stages"] = settings["stages"]
        if data:
            scenario.production = data
            scenario.save()
            prewarm(scenario)
            messages.success(request, "Production plan saved and recalculated.")
            return redirect("production", pk=pk)
        for e in errors:
            messages.error(request, e)
    ctx = {"scenario": scenario, "scenarios": Scenario.objects.all(), "settings": settings,
           "panels": _stage_panels(settings, base)}
    got, err = plan_result(scenario)
    if err:
        ctx["error"] = err
        return render(request, "planner/production.html", ctx)
    plan, res = got

    chart = {
        "start": plan.start_date.isoformat(), "horizon": res["horizon"],
        "stages": [s.name for s in res["stages"]],
        "volume": [[round(d["volume_m3"], 2) for d in days] for days in res["daily"]],
        "feed": [[round(d["feed_kg"], 2) for d in days] for days in res["daily"]],
        "batches": [[d["batches"] for d in days] for days in res["daily"]],
        "gantt": [{"label": f"B{b['batch']}", "segments": [
            {"stage": s["stage"], "start": s["start"], "end": s["end"] + 1,
             "w_in": round(s["w_in"], 2), "w_out": round(s["w_out"], 2)} for s in b["segments"]],
            "harvest": b["harvest_day"]} for b in res["batches"]],
    }
    total_volume = sum(d["peak"]["volume_m3"] for d in res["designs"] if not d["empty"]) or 1
    ctx.update({"res": res, "plan": plan, "chart": chart, "designs": res["designs"],
                "total_volume": total_volume,
                "harvests": res["harvests"],
                "fish_per_batch": scenario.engine_params().growth.initial_count})
    return render(request, "planner/production.html", ctx)


def standards(request):
    edit = get_object_or_404(EquipmentStandard, pk=request.GET["edit"]) if "edit" in request.GET else None
    if request.method == "POST":
        if "delete" in request.POST:
            EquipmentStandard.objects.filter(pk=request.POST["delete"]).delete()
            messages.success(request, "Standard deleted.")
            return redirect("standards")
        instance = (get_object_or_404(EquipmentStandard, pk=request.POST["id"])
                    if request.POST.get("id") else None)
        form = StandardForm(request.POST, instance=instance)
        if form.is_valid():
            form.save()
            messages.success(request, "Standard saved.")
            return redirect("standards")
    else:
        initial = {"category": request.GET.get("category")} if "category" in request.GET else None
        form = StandardForm(instance=edit, initial=initial)
    groups = []
    for key, label, unit in CATEGORIES:
        items = [s for s in EquipmentStandard.objects.all() if s.category == key]
        groups.append({"key": key, "label": label, "unit": unit, "items": items})
    return render(request, "planner/standards.html", {
        "groups": groups, "form": form, "edit": edit,
        "scenario": Scenario.objects.first(), "scenarios": Scenario.objects.all()})


def equipment(request, pk):
    scenario = get_object_or_404(Scenario, pk=pk)
    saved = scenario.equipment or {}
    if request.method == "POST":
        stages_sel = {}
        n = _int_or_none(request.POST.get("n_stages")) or 0
        for i in range(n):
            sel = {}
            for key, _label, _unit in SIZED_CATEGORIES:
                p = f"s{i}-{key}-"
                sel[key] = {"item": request.POST.get(p + "item", ""),
                            "capacity": _float_or_none(request.POST.get(p + "cap", "")),
                            "qty": _int_or_none(request.POST.get(p + "qty", ""))}
            # tanks per batch by weight: "up to X g -> N tanks", last step open-ended
            steps = [[_float_or_none(u), _int_or_none(k)] for u, k in
                     zip(request.POST.getlist(f"s{i}-tank-upto"), request.POST.getlist(f"s{i}-tank-k"))]
            if f"s{i}-tank-item" in request.POST:
                sel["tank"]["steps"] = [list(s) for s in normalize_steps(steps)]
                sel["tank"]["qty"] = None
            else:  # tanks are chosen on the Units & tanks step
                sel["tank"] = scenario.equipment_of(i).get("tank", {})
            stages_sel[scenario.unit_id(i)] = sel
            _save_diagram_picks(request.POST, scenario, i)
        scenario.equipment = {"water": request.POST.get("water", "fresh"), "stages": stages_sel}
        scenario.save()
        messages.success(request, "Equipment saved and checked against the plan.")
        return redirect("equipment", pk=pk)

    ctx = {"scenario": scenario, "scenarios": Scenario.objects.all()}
    got, err = plan_result(scenario)
    if err:
        ctx["error"] = err
        return render(request, "planner/equipment.html", ctx)
    plan, res = got

    units, occupancy, water = _equipment_units(scenario, res)
    ctx.update({"plan": plan, "res": res, "units": units, "water": water,
                "grand_total": sum(u["total"] for u in units),
                "any_price": any(u["priced"] or u["diagram_priced"] for u in units),
                "occupancy": {"start": plan.start_date.isoformat(), "horizon": res["horizon"],
                              "stages": occupancy}})
    return render(request, "planner/equipment.html", ctx)


def _save_diagram_picks(post, scenario, index):
    """Standard and unit price picked on Equipment for a group of identical diagram items: written to each of
    those items in the diagram, so the diagram stays the one place that says what is installed."""
    diagram = scenario.diagram_of(index)
    groups = [k[len(f"d{index}-"):-len("-nodes")] for k in post if k.startswith(f"d{index}-") and k.endswith("-nodes")]
    if not diagram or not groups:
        return
    picks = {}
    for g in groups:
        p = f"d{index}-{g}-"
        std = post.get(p + "std", "")
        price = _float_or_none(post.get(p + "price", ""))
        for nid in post.get(p + "nodes", "").split(","):
            picks[nid] = (std if std.isdigit() else "", "" if price is None else price)
    changed = False
    for n in diagram.get("nodes") or []:
        if n.get("id") in picks:
            data = n.setdefault("data", {})
            std, price = picks[n["id"]]
            if (str(data.get("standard") or ""), data.get("unit_price") if data.get("unit_price") is not None else "") != (std, price):
                data["standard"], data["unit_price"] = std, price
                changed = True
    if changed:
        scenario.set_diagram(index, diagram)


def _library_list():
    return [{"id": t.pk, "name": t.name, "source": t.source, "updated": t.updated.strftime("%d %b %Y %H:%M"),
             "items": sum(1 for n in t.diagram.get("nodes", []) if NODE_TYPES.get(n.get("type"), {}).get("category"))}
            for t in ProcessTemplate.objects.all()]


def process_library_item(request, tid):
    """GET: one saved diagram. POST: delete it."""
    t = get_object_or_404(ProcessTemplate, pk=tid)
    if request.method == "POST":
        t.delete()
        return JsonResponse({"ok": True, "library": _library_list()})
    return JsonResponse({"ok": True, "name": t.name, "diagram": t.diagram})


CAPEX_GROUPS = ["Building & civil works", "Piping & valves", "Electrical & control", "Installation",
                "Engineering & design", "Transport", "Contingency", "Other"]


def _capex_extras_from_post(post):
    """The 'other investment items' rows exactly as typed (step 07), for saving or re-rendering."""
    extras = []
    for name, group, mode, qty, price, pct in zip(
            post.getlist("x_name"), post.getlist("x_group"), post.getlist("x_mode"),
            post.getlist("x_qty"), post.getlist("x_price"), post.getlist("x_pct")):
        if not name.strip():
            continue
        extras.append({"name": name.strip(), "group": group, "mode": mode,
                       "qty": _float_or_none(qty), "unit_price": _float_or_none(price),
                       "percent": _float_or_none(pct)})
    return extras


def _capex_page(request, scenario, post=None, status=200):
    """Step 07 rendered. With `post` (a save refused as stale), the typed prices and other items
    show what was typed; everything else is the page as saved."""
    saved = scenario.capex or {}
    currency = (post.get("currency", "") if post is not None else saved.get("currency", "EUR"))
    ctx = {"scenario": scenario, "scenarios": Scenario.objects.all(), "groups": CAPEX_GROUPS,
           "currency": currency.strip() or "EUR"}
    got, err = plan_result(scenario)
    if err:
        ctx["error"] = err
        return render(request, "planner/capex.html", ctx, status=status)
    _plan, res = got
    ctx.update(_capex_data(scenario, res))
    if post is not None:
        for ln in ctx["lines"]:  # the typed prices stand in the inputs; totals stay as saved
            price = _float_or_none(post.get(f"price-{ln['key']}", ""))
            if price is not None:
                ln.update(price=price, overridden=True)
        ctx["extras"] = [{**x, "total": None} for x in _capex_extras_from_post(post)]
    return render(request, "planner/capex.html", ctx, status=status)


def capex(request, pk):
    scenario = get_object_or_404(Scenario, pk=pk)
    if request.method == "POST":
        # optimistic concurrency: a mismatching `updated` means another tab saved since this form
        # was rendered — show what was typed over the current page, save nothing
        posted_updated = request.POST.get("updated", "").strip()
        if posted_updated and posted_updated != scenario.updated.isoformat():
            messages.warning(request, f"This scenario was saved by another tab at {scenario.updated}. "
                                      "Reload the page and re-apply your changes.")
            return _capex_page(request, scenario, post=request.POST, status=409)
        overrides = {}
        for key in request.POST.getlist("line_key"):
            price = _float_or_none(request.POST.get(f"price-{key}", ""))
            if price is not None:
                overrides[key] = price
        extras = _capex_extras_from_post(request.POST)
        scenario.capex = {"currency": request.POST.get("currency", "EUR").strip() or "EUR",
                          "price_overrides": overrides, "extras": extras}
        # prune overrides that no longer match any CAPEX line (its key changes only when the drawn
        # equipment itself changes — but leftovers from older saves must not linger silently)
        got, err = plan_result(scenario)
        if not err:
            live_keys = {ln["key"] for ln in _capex_data(scenario, got[1])["lines"]}
            stale = [k for k in overrides if k not in live_keys]
            if stale:
                scenario.capex["price_overrides"] = {k: v for k, v in overrides.items() if k in live_keys}
                messages.warning(request, f"{len(stale)} old price override(s) no longer match a CAPEX line "
                                           "and were removed")
        scenario.save()
        messages.success(request, "CAPEX saved.")
        return redirect("capex", pk=pk)

    return _capex_page(request, scenario)


def feeds(request):
    edit = get_object_or_404(Feed, pk=request.GET["edit"]) if "edit" in request.GET else None
    if request.method == "POST":
        if "delete" in request.POST:
            Feed.objects.filter(pk=request.POST["delete"]).delete()
            messages.success(request, "Feed deleted.")
            return redirect("feeds")
        instance = get_object_or_404(Feed, pk=request.POST["id"]) if request.POST.get("id") else None
        form = FeedForm(request.POST, instance=instance)
        if form.is_valid():
            form.save()
            messages.success(request, "Feed saved.")
            return redirect("feeds")
    else:
        form = FeedForm(instance=edit)
    catalog = list(Feed.objects.all())
    gaps = [(a.max_weight_g, b.min_weight_g) for a, b in zip(catalog, catalog[1:])
            if b.min_weight_g != a.max_weight_g]
    return render(request, "planner/feeds.html", {
        "feeds": catalog, "form": form, "edit": edit, "gaps": gaps,
        "scenario": Scenario.objects.first(), "scenarios": Scenario.objects.all()})


def export_xlsx(request, pk):
    scenario = get_object_or_404(Scenario, pk=pk)
    try:
        rows = _run(scenario)
    except ValueError as exc:  # e.g. the feed table does not cover the weight range
        messages.error(request, str(exc))
        return redirect("dashboard", pk=pk)
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily"
    ws.append([label for _, label in EXPORT_COLUMNS])
    for r in rows:
        ws.append([_finite(r.get(k), 6) if isinstance(r.get(k), float) else r.get(k)
                   for k, _ in EXPORT_COLUMNS])
    ws2 = wb.create_sheet("Weekly")
    ws2.append([label for _, label in EXPORT_COLUMNS])
    for r in weekly(rows):
        ws2.append([_finite(r.get(k), 6) if isinstance(r.get(k), float) else r.get(k)
                    for k, _ in EXPORT_COLUMNS])
    ws3 = wb.create_sheet("Constants")
    for k, v in scenario.engine_params().to_flat().items():
        ws3.append([k, v])
    ws3.append(["temperature_steps", scenario.temperature_steps])
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type=(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
    resp["Content-Disposition"] = f'attachment; filename="production_plan_{scenario.pk}.xlsx"'
    return resp
