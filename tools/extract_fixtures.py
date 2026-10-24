"""Extract the values trout.xlsx computed into JSON as test references.

Usage:  python tools/extract_fixtures.py trout.xlsx [--manifest]

--manifest additionally writes tests/fixtures/MANIFEST.txt with the SHA-256
hashes of the source workbook and the generated fixtures.
"""
import hashlib
import json
import subprocess
import sys
import warnings
from datetime import datetime
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")

DFRY_COLUMNS = {
    # inputs
    "F": "day", "G": "temperature", "D": "fcr_correction", "S": "protein",
    "AH": "salinity", "AI": "cone_factor",
    # outputs
    "E": "sgr", "H": "length_inc", "I": "weight_g", "J": "length_cm", "K": "count",
    "L": "biomass_kg", "M": "biomass_t", "N": "density", "O": "volume_m3", "P": "fcr",
    "Q": "feed_kg", "R": "feed_rate", "T": "feed_cum_t",
    "U": "o2_lpm", "V": "o2_lpm_actual", "W": "o2_lph", "X": "o2_kgh", "Y": "o2_kgd",
    "Z": "co2_lpm", "AA": "co2_lph", "AB": "co2_kgh", "AC": "co2_kgd",
    "AE": "ozone_kgd", "AF": "ozone_gh", "AG": "ozone_lpm",
    "AJ": "o2_sat_in", "AK": "o2_sat_out",
    "AL": "q_o2_m3h", "AM": "q_o2_lpm", "AN": "q_tan_m3h", "AO": "q_tan_lpm",
    "AP": "q_co2_m3h", "AQ": "q_co2_lpm", "AR": "q_tss_m3h", "AS": "q_tss_lpm",
    "AT": "q_no3_m3d", "AU": "water_per_feed", "AV": "backflush60_m3d", "AW": "backflush30_m3d",
}

BIOFILTER_COLUMNS = {
    "E": "retention_min", "F": "q_tank_m3h", "H": "q_biofilter_m3h", "I": "tan_out",
    "J": "efficiency_pct", "L": "monod_rate", "M": "k_temp", "N": "k_media",
    "O": "rate_salt", "P": "rate_fresh", "Q": "media_salt_m3", "R": "media_fresh_m3",
    "S": "volume_salt_m3", "T": "volume_fresh_m3", "U": "residence_salt_min",
    "V": "residence_fresh_min",
}


def num(v):
    return v if isinstance(v, (int, float)) else None


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_manifest(path, out_dir):
    root = Path(__file__).resolve().parent.parent
    try:
        head = subprocess.run(
            ["git", "rev-parse", "HEAD"], cwd=root, capture_output=True,
            text=True, check=True).stdout.strip()
    except Exception:
        head = "(unavailable)"
    lines = [
        f"source_workbook: {Path(path).name}",
        f"source_sha256: {sha256(path)}",
        f"dfry.json_sha256: {sha256(out_dir / 'dfry.json')}",
        f"biofilter.json_sha256: {sha256(out_dir / 'biofilter.json')}",
        f"generated_at: {datetime.now().isoformat(timespec='seconds')}",
        f"git_head: {head}",
    ]
    manifest = out_dir / "MANIFEST.txt"
    manifest.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"manifest -> {manifest}")


def main(path, manifest=False):
    wb = openpyxl.load_workbook(path, data_only=True)
    out_dir = Path(__file__).resolve().parent.parent / "tests" / "fixtures"
    out_dir.mkdir(parents=True, exist_ok=True)

    ws = wb["D. Fry"]
    rows = []
    for r in range(9, 737):
        rows.append({name: num(ws[f"{col}{r}"].value) for col, name in DFRY_COLUMNS.items()})
    (out_dir / "dfry.json").write_text(json.dumps(rows, indent=0), encoding="utf-8")

    ws = wb["Sayfa11"]
    inputs = {
        "feeding_hours": ws["C7"].value, "temperature": ws["D8"].value,
        "tank_volume_m3": ws["D14"].value, "n_tanks": ws["D15"].value,
        "protein": ws["D21"].value, "daily_feed_kg": ws["D22"].value,
        "media_area_m2m3": ws["D28"].value, "tan_target": ws["D29"].value,
        "makeup_flow_m3h": ws["D31"].value or 0, "tan_inlet": ws["D33"].value,
        "fill_ratio_pct": ws["R10"].value,
    }
    bio_rows = [
        {name: num(ws[f"{col}{r}"].value) for col, name in BIOFILTER_COLUMNS.items()}
        for r in range(13, 34)
    ]
    (out_dir / "biofilter.json").write_text(
        json.dumps({"inputs": inputs, "rows": bio_rows}, indent=0), encoding="utf-8")
    print(f"{len(rows)} days, {len(bio_rows)} biofilter rows -> {out_dir}")

    if manifest:
        write_manifest(path, out_dir)


if __name__ == "__main__":
    argv = [a for a in sys.argv[1:] if a != "--manifest"]
    main(argv[0] if argv else "trout.xlsx", manifest="--manifest" in sys.argv)
